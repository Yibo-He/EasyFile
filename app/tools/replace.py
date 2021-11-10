from docx.shared import Pt, RGBColor

from . import runtools
from .myfont import Font
from .myfont import RepStr


def get_typeface(run, paragraph, default):
    name = default.name
    if run.font.name is not None:
        name = run.font.name
    else:
        style = run.style
        while style is not None and style.font.name is None:
            style = style.base_style
        if style is not None:
            name = style.font.name
        elif paragraph.style.font.name != None:
            name = paragraph.style.font.name
        else:
            style = paragraph.style
            while style is not None and style.font.name is None:
                style = style.base_style
            if style is not None:
                name = style.font.name
    return name


def get_size(run, paragraph, default):
    size = default.size
    if run.font.size is not None:
        size = run.font.size.pt
    else:
        style = run.style
        while style is not None and style.font.size is None:
            style = style.base_style
        if style is not None:
            size = style.font.size
        elif paragraph.style.font.size is not None:
            size = paragraph.style.font.size
        else:
            style = paragraph.style
            while style is not None and style.font.size is None:
                style = style.base_style
            if style is not None:
                size = style.font.size
    return size


def get_rgbcolor(run, paragraph, default):
    color = default.color
    if run.font.color.type is not None:
        color = run.font.color.rgb
    else:
        style = run.style
        while style is not None and style.font.color.rgb is None:
            style = style.base_style
        if style is not None:
            color = style.font.color.rgb
        elif paragraph.style.font.color.rgb is not None:
            color = paragraph.style.font.color.rgb
        else:
            style = paragraph.style
            while style is not None and style.font.color.rgb is None:
                style = style.base_style
            if style is not None:
                color = style.font.color.rgb
    return color


def get_font(run, paragraph, default):
    name = get_typeface(run, paragraph, default)
    color = get_rgbcolor(run, paragraph, default)
    size = get_size(run, paragraph, default)
    return Font(color=color, name=name, size=size)


def check(old: RepStr, now: RepStr):
    ret = []
    if(now.str[0]=='用'):
        t=1
    ret.append(old.str == "" or old.str == now.str)
    ret.append(old.font.color == "" or old.font.color == now.font.color)
    ret.append(old.font.size == 0 or old.font.size == now.font.size)
    #   if(now.font.size==16.0):
    #       print(now.str)
    ret.append(old.font.name == "" or old.font.name == now.font.name)
    for x in ret:
        if x is False:
            return False
    #   print(now.font.name)
    return True


def replace_char(document, default, old, new):
    for para in document.paragraphs:
        for run in para.runs:
            if (len(run.text) > 1):
                t = range(len(run.text))
                runtools.split_run_by(para, run, t)
    myruns = []
    for para in document.paragraphs:
        for run in para.runs:
            myruns.append(run)

    for para in document.paragraphs:
        for i, run in enumerate(para.runs):
            nowfont = get_font(run, para, default)
            # print(run.text)
            if (len(run.text) > 1):
                t = range(len(run.text))
                ret = runtools.split_run_by(para, run, t)
            for j, x in enumerate(run.text):
                if check(old, RepStr(x, nowfont.color, nowfont.name, nowfont.size)):
                    lst = list(para.runs[i].text)
                    lst[j] = new.str[0]
                    run.text = "".join(lst)
                    run.font.color.rgb = RGBColor.from_string(new.font.color)
                    # print(new.font.name)
                    run.font.name = new.font.name
                    run.font.size = Pt(new.font.size)
    return


def replace(document, default, old, new):
    if len(old.str) < 2:
        replace_char(document, default, old, new)
        return

    for para in document.paragraphs:
        for run in para.runs:
            if (len(run.text) > 1):
                t = range(len(run.text))
                runtools.split_run_by(para, run, t)
    myruns = []
    for para in document.paragraphs:
        for run in para.runs:
            if len(run.text) < 1:
                continue
            nowfont = get_font(run, para, default)
            myruns.append([run, nowfont])
    print(old.str)

    t = RepStr.from_font(old.str[0], old.font)
    print(t.str)
    print(t.font.name)
    print(t.font.size)
    print(t.font.color)
    for i in range(len(myruns) - len(old.str)):
        flag = 0
        if len(myruns[i][0].text) ==0 :
            continue
        for j in range(len(old.str)):
            assert(len(myruns[i + j][0].text)>0)
            if check(RepStr(old.str[j], old.font.color,old.font.name,old.font.size), RepStr(myruns[i + j][0].text[0], myruns[i + j][1].color,myruns[i + j][1].name,myruns[i + j][1].size)) is False:
                flag = 1
                break
        if not flag:
            run = myruns[i][0]
            run.text = new.str
            run.font.color.rgb = RGBColor.from_string(new.font.color)
            # print(new.font.name)
            run.font.name = new.font.name
            run.font.size = Pt(new.font.size)
            for j in range(1, len(old.str)):
                myruns[i + j][0].text = ""
    return
